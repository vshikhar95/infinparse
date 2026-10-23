"""Zerodha Console holdings export (holdings-<client id>.xlsx). Needs the `excel` extra (openpyxl)."""
import re
from datetime import date
from decimal import Decimal
from pathlib import Path

from indian_finance_parsers.models import Holding

QTY = ["Quantity Available", "Quantity Discrepant", "Quantity Pledged (Margin)", "Quantity Pledged (Loan)"]


def _dec(v) -> Decimal:
    return Decimal(str(v or 0).replace(",", "").strip() or 0)


def parse(path, as_of: date | None = None) -> list[Holding]:
    """Rows under the Symbol/ISIN header until the first blank Symbol. From the preamble: account = the Client ID,
    or ""; as_of = the "Statement as on YYYY-MM-DD" date, else the as_of argument, else the file's mtime."""
    from openpyxl import load_workbook
    rows = iter(load_workbook(path, data_only=True).active.iter_rows(values_only=True))
    client, printed, totals = "", None, {}
    for hdr in rows:
        cells = [c for c in hdr or () if c is not None]  # the export starts in column B, so find labels anywhere
        if "Client ID" in cells[:-1]:
            client = str(cells[cells.index("Client ID") + 1])
        if m := next((re.search(r"as on (\d{4}-\d\d-\d\d)", c) for c in cells if isinstance(c, str) and "as on" in c), None):
            printed = date.fromisoformat(m[1])
        for lab in ("Invested Value", "Present Value"):  # older exports don't print them
            if lab in cells[:-1]:
                totals[lab] = _dec(cells[cells.index(lab) + 1])
        if "Symbol" in cells and "ISIN" in cells:
            break
    else:
        raise ValueError(f"{Path(getattr(path, 'name', path)).name}: no Symbol/ISIN header row")
    if not (printed or as_of or isinstance(path, (str, Path))):
        raise ValueError(f"{getattr(path, 'name', 'file')}: no statement date in the file; pass as_of")
    as_of = printed or as_of or date.fromtimestamp(Path(path).stat().st_mtime)
    ix = {h: i for i, h in enumerate(hdr)}
    out, invested, present = [], Decimal(0), Decimal(0)
    for r in rows:
        if not r[ix["Symbol"]]:
            break
        units = sum(_dec(r[ix[q]]) for q in QTY)
        price = _dec(r[ix["Previous Closing Price"]])
        avg = _dec(r[ix["Average Price"]]) if "Average Price" in ix else Decimal(0)
        isin = r[ix["ISIN"]]
        if not isin and re.fullmatch(r"IN[A-Z0-9]{10}", str(r[ix["Symbol"]])):
            isin = r[ix["Symbol"]]  # an unlisted bond (54EC) is listed under its ISIN with the ISIN cell blank
        invested += units * avg
        present += units * price
        out.append(Holding(as_of, "zerodha", client, r[ix["Symbol"]], isin, units, price, units * price,
                           cost=units * avg if avg else None))  # the sheet's average buy price x the units held
    name = Path(getattr(path, "name", path)).name  # the sheet's own totals: a dropped or misread row rejects the file
    for field, printed_total, parsed in (("Invested Value", totals.get("Invested Value"), invested if "Average Price" in ix else None),
                                         ("Present Value", totals.get("Present Value"), present)):
        if printed_total is not None and parsed is not None and abs(parsed - printed_total) > 1:  # real sheets are ₹0.01 off
            raise ValueError(f"{name}: parsed {field} doesn't match the sheet's {field} (a row was missed or misread)")
    return out
