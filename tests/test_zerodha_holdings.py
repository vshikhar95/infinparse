import os
from datetime import date
from decimal import Decimal as D

import pytest
from openpyxl import Workbook

from indian_finance_parsers.portfolio import zerodha_holdings

HDR = ["Symbol", "ISIN", "Sector", "Quantity Available", "Quantity Discrepant", "Quantity Long Term",
       "Quantity Pledged (Margin)", "Quantity Pledged (Loan)", "Average Price", "Previous Closing Price",
       "Unrealized P&L", "Unrealized P&L Pct."]


def _holdings_xlsx(tmp_path):
    p = tmp_path / "holdings-AB1234.xlsx"
    wb = Workbook()
    for r in [["Equity Holdings Statement as on 2026-10-20"], [], ["Client ID", "AB1234"], [], HDR,
              ["FAKECO", "INE000X01011", "IT", 10, 1, 5, 2, 0, 100, 150.5, 0, 0],
              ["TESTBANK", "INE000X01029", "BANKING", 3, 0, 0, 0, 1, 200, 210, 0, 0],
              [None, None, "Total", 16],
              ["Notes: invented"]]:
        wb.active.append(r)
    wb.save(p)
    return p


def test_holdings(tmp_path):
    rows = zerodha_holdings.parse(_holdings_xlsx(tmp_path), as_of=date(2026, 10, 20))
    assert [(r.name, r.isin, r.units, r.price, r.value) for r in rows] == [
        ("FAKECO", "INE000X01011", D(13), D("150.5"), D("1956.5")),
        ("TESTBANK", "INE000X01029", D(4), D(210), D(840))]
    assert (rows[0].source, rows[0].as_of, rows[0].matures) == ("zerodha", date(2026, 10, 20), None)


def test_account_is_the_client_id_printed_in_the_file(tmp_path):
    assert {r.account for r in zerodha_holdings.parse(_holdings_xlsx(tmp_path))} == {"AB1234"}


def test_account_is_empty_when_the_file_prints_none(tmp_path):
    from openpyxl import load_workbook
    p = _holdings_xlsx(tmp_path)
    wb = load_workbook(p)
    wb.active.delete_rows(3)
    wb.save(p)
    assert {r.account for r in zerodha_holdings.parse(p)} == {""}


def test_as_of_defaults_to_mtime(tmp_path):
    p = _holdings_xlsx(tmp_path)
    os.utime(p, (1_800_000_000, 1_800_000_000))
    assert zerodha_holdings.parse(p)[0].as_of == date.fromtimestamp(1_800_000_000)


def test_no_header_raises_naming_only_the_file(tmp_path):
    p = tmp_path / "holdings.xlsx"
    wb = Workbook()
    wb.active.append(["nothing here"])
    wb.save(p)
    with pytest.raises(ValueError, match="^holdings.xlsx: no Symbol/ISIN header row"):
        zerodha_holdings.parse(p)
