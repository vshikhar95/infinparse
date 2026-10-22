"""Zerodha Console holdings export (holdings-<client id>.xlsx). Needs the `excel` extra (openpyxl)."""
from datetime import date
from decimal import Decimal
from pathlib import Path

from indian_finance_parsers.models import Holding

QTY = ["Quantity Available", "Quantity Discrepant", "Quantity Pledged (Margin)", "Quantity Pledged (Loan)"]


def _dec(v) -> Decimal:
    return Decimal(str(v or 0).replace(",", "").strip() or 0)


def parse(path, as_of: date | None = None) -> list[Holding]:
    """Rows under the Symbol/ISIN header until the first blank Symbol; as_of defaults to the file's mtime
    (the export carries no valuation date of its own). account = the Client ID in the preamble, or ""."""
    from openpyxl import load_workbook
    as_of = as_of or date.fromtimestamp(Path(path).stat().st_mtime)
    rows = iter(load_workbook(path, data_only=True).active.iter_rows(values_only=True))
    client = ""
    for hdr in rows:
        if hdr and hdr[0] == "Client ID":
            client = str(hdr[1] or "")
        if hdr and "Symbol" in hdr and "ISIN" in hdr:
            break
    else:
        raise ValueError(f"{Path(path).name}: no Symbol/ISIN header row")
    ix = {h: i for i, h in enumerate(hdr)}
    out = []
    for r in rows:
        if not r[ix["Symbol"]]:
            break
        units = sum(_dec(r[ix[q]]) for q in QTY)
        price = _dec(r[ix["Previous Closing Price"]])
        out.append(Holding(as_of, "zerodha", client, r[ix["Symbol"]], r[ix["ISIN"]], units, price, units * price))
    return out
