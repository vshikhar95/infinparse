"""Zerodha holdings xlsx, tradebook CSV and ledger CSV importers."""
import csv
from datetime import date
from decimal import Decimal
from pathlib import Path

from indian_finance_parsers.models import Holding, Trade, Txn

QTY = ["Quantity Available", "Quantity Discrepant", "Quantity Pledged (Margin)", "Quantity Pledged (Loan)"]


def _dec(v) -> Decimal:
    return Decimal(str(v or 0).replace(",", "").strip() or 0)


def holdings(path, as_of: date | None = None, owner: str = "own") -> list[Holding]:
    """Rows under the Symbol/ISIN header until the first blank Symbol; as_of defaults to file mtime."""
    from openpyxl import load_workbook  # the excel extra; trades() and ledger() don't need it
    as_of = as_of or date.fromtimestamp(Path(path).stat().st_mtime)
    rows = iter(load_workbook(path, data_only=True).active.iter_rows(values_only=True))
    for hdr in rows:
        if hdr and "Symbol" in hdr and "ISIN" in hdr:
            break
    else:
        raise ValueError(f"{path}: no Symbol/ISIN header row")
    ix = {h: i for i, h in enumerate(hdr)}
    out = []
    for r in rows:
        if not r[ix["Symbol"]]:
            break
        units = sum(_dec(r[ix[q]]) for q in QTY)
        price = _dec(r[ix["Previous Closing Price"]])
        out.append(Holding(as_of, "zerodha", "Zerodha", r[ix["Symbol"]], r[ix["ISIN"]], units, price, units * price, owner))
    return out


def trades(path) -> list[Trade]:
    with open(path, newline="") as f:
        rows = list(csv.DictReader(f))
    bad = {r["trade_type"] for r in rows} - {"buy", "sell"}
    if bad:
        raise ValueError(f"{path}: unexpected trade_type {sorted(bad)}")
    return [Trade(r["trade_id"], date.fromisoformat(r["trade_date"]), r["symbol"], r["isin"], r["trade_type"],
                      Decimal(r["quantity"]), Decimal(r["price"]), r["exchange"]) for r in rows]


def _ledger_rows(path):
    with open(path, newline="") as f:
        return list(csv.DictReader(f))


def opening(path) -> Decimal:
    return _dec(_ledger_rows(path)[0]["net_balance"])


def closing(path) -> Decimal | None:
    """Net balance of the trailing Closing Balance row, if the file has one."""
    last = _ledger_rows(path)[-1]
    return _dec(last["net_balance"]) if last["particulars"] == "Closing Balance" else None


def ledger(path) -> list[Txn]:
    """Dated rows only (Opening/Closing Balance skipped); raises if Closing Balance disagrees with the last row."""
    txns = [Txn(date.fromisoformat(r["posting_date"]), _dec(r["credit"]) - _dec(r["debit"]), r["particulars"],
                _dec(r["net_balance"])) for r in _ledger_rows(path) if r["posting_date"]]
    end = closing(path)
    last = txns[-1].balance_after if txns else opening(path)
    if end is not None and end != last:
        raise ValueError(f"{path}: closing balance {end} != last row balance {last}")
    return txns
